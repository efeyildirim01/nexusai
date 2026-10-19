"""PDF and Excel exports of a dashboard view."""

import io
from datetime import UTC, datetime
from typing import Any

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.barcharts import HorizontalBarChart
from reportlab.graphics.charts.lineplots import LinePlot
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.analytics import Filters
from app.models import SalesRecord as S

MAX_EXPORT_ROWS = 100_000
BRAND = colors.HexColor("#4F46E5")
MUTED = colors.HexColor("#6B7280")


def _fmt_money(value: float) -> str:
    return f"{value:,.2f}"


def _fmt_delta(value: float | None) -> str:
    return "n/a" if value is None else f"{value:+.1f}%"


def _filters_line(filters: Filters, overview: dict[str, Any]) -> str:
    parts = [f"{overview['period']['date_from']} to {overview['period']['date_to']}"]
    for label, value in (
        ("Region", filters.region),
        ("Category", filters.category),
        ("Channel", filters.channel),
    ):
        if value:
            parts.append(f"{label}: {value}")
    return " · ".join(parts)


def _line_chart(series: list[dict[str, Any]], width: float, height: float) -> Drawing:
    drawing = Drawing(width, height)
    if len(series) < 2:
        drawing.add(String(10, height / 2, "Not enough data for a trend", fillColor=MUTED))
        return drawing
    plot = LinePlot()
    plot.x, plot.y = 45, 20
    plot.width, plot.height = int(width - 60), int(height - 35)
    plot.data = [[(i, row["revenue"]) for i, row in enumerate(series)]]
    plot.lines[0].strokeColor = BRAND
    plot.lines[0].strokeWidth = 1.6
    plot.xValueAxis.valueMin = 0
    plot.xValueAxis.valueMax = len(series) - 1
    step = max(1, len(series) // 6)
    plot.xValueAxis.valueSteps = list(range(0, len(series), step))
    plot.xValueAxis.labelTextFormat = lambda i: series[int(i)]["period"][2:]
    plot.xValueAxis.labels.fontSize = 6
    plot.yValueAxis.valueMin = 0
    plot.yValueAxis.labels.fontSize = 6
    plot.yValueAxis.labelTextFormat = lambda v: f"{v / 1000:,.0f}k" if v >= 1000 else f"{v:,.0f}"
    drawing.add(plot)
    return drawing


def _bar_chart(rows: list[dict[str, Any]], width: float, height: float) -> Drawing:
    drawing = Drawing(width, height)
    rows = list(reversed(rows[:6]))
    if not rows:
        return drawing
    chart = HorizontalBarChart()
    chart.x, chart.y = 90, 10
    chart.width, chart.height = int(width - 110), int(height - 20)
    chart.data = [[r["revenue"] for r in rows]]
    chart.categoryAxis.categoryNames = [str(r["key"])[:18] for r in rows]
    chart.categoryAxis.labels.fontSize = 7
    chart.valueAxis.valueMin = 0
    chart.valueAxis.labels.fontSize = 6
    chart.bars[0].fillColor = BRAND
    chart.bars[0].strokeColor = None
    drawing.add(chart)
    return drawing


def _table(header: list[str], rows: list[list[str]], widths: list[float]) -> Table:
    table = Table([header, *rows], colWidths=widths, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
                ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
                ("TEXTCOLOR", (0, 0), (-1, 0), MUTED),
                ("LINEBELOW", (0, 0), (-1, 0), 0.5, MUTED),
                ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F5F7")]),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def build_pdf(
    tenant_name: str, filters: Filters, overview: dict[str, Any], narrative: dict[str, Any]
) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"{tenant_name} sales report",
        author="NexusAI",
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], alignment=0, fontSize=18, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=11, spaceBefore=10, spaceAfter=4)
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9, leading=12)
    small = ParagraphStyle("small", parent=body, fontSize=8, textColor=MUTED)
    width = A4[0] - 36 * mm

    k, d = overview["kpis"], overview["deltas_pct"]
    story: list[Any] = [
        Paragraph(f"{tenant_name} — Sales report", h1),
        Paragraph(_filters_line(filters, overview), small),
        Paragraph(f"Generated {datetime.now(UTC):%Y-%m-%d %H:%M} UTC by NexusAI", small),
        Spacer(1, 6),
        Paragraph(f"<b>{narrative['headline']}</b>", body),
        Paragraph(narrative["summary"], body),
        Paragraph("Key figures", h2),
        _table(
            ["Metric", "Value", "vs previous period"],
            [
                ["Revenue", _fmt_money(k["revenue"]), _fmt_delta(d["revenue"])],
                ["Orders", f"{k['orders']:,}", _fmt_delta(d["orders"])],
                ["Customers", f"{k['customers']:,}", _fmt_delta(d["customers"])],
                [
                    "Average order value",
                    _fmt_money(k["average_order_value"]),
                    _fmt_delta(d["average_order_value"]),
                ],
                ["Gross margin", f"{k['gross_margin_pct']:.1f}%", ""],
                ["Return rate", f"{k['return_rate_pct']:.1f}%", ""],
            ],
            [width * 0.4, width * 0.3, width * 0.3],
        ),
        Paragraph(f"Revenue by {overview['period']['grain']}", h2),
        _line_chart(overview["timeseries"], width, 150),
        Paragraph("Revenue by category", h2),
        _bar_chart(overview["by_category"], width, 120),
        Paragraph("Top products", h2),
        _table(
            ["Product", "Revenue", "Lines"],
            [
                [r["key"], _fmt_money(r["revenue"]), f"{r['lines']:,}"]
                for r in overview["top_products"]
            ],
            [width * 0.5, width * 0.3, width * 0.2],
        ),
        Paragraph("Top regions", h2),
        _table(
            ["Region", "Revenue", "Lines"],
            [
                [r["key"], _fmt_money(r["revenue"]), f"{r['lines']:,}"]
                for r in overview["by_region"]
            ],
            [width * 0.5, width * 0.3, width * 0.2],
        ),
    ]
    if overview["anomalies"]:
        story.append(Paragraph("Anomalies", h2))
        for anomaly, text in zip(
            overview["anomalies"], narrative["anomaly_explanations"], strict=False
        ):
            story.append(Paragraph(f"<b>{anomaly['period']}</b> — {text}", body))
    doc.build(story)
    return buffer.getvalue()


def _sheet(writer: pd.ExcelWriter, name: str, df: pd.DataFrame) -> None:
    df.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    header_fill = PatternFill("solid", fgColor="4F46E5")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    for i, column in enumerate(df.columns, start=1):
        values = [str(column), *(str(v) for v in df[column].head(200))]
        ws.column_dimensions[get_column_letter(i)].width = min(45, max(10, *map(len, values)) + 2)
    ws.freeze_panes = "A2"


def build_xlsx(session: Session, filters: Filters, overview: dict[str, Any]) -> bytes:
    k, d = overview["kpis"], overview["deltas_pct"]
    summary = pd.DataFrame(
        [
            ("Period from", overview["period"]["date_from"], None),
            ("Period to", overview["period"]["date_to"], None),
            ("Revenue", k["revenue"], d["revenue"]),
            ("Orders", k["orders"], d["orders"]),
            ("Customers", k["customers"], d["customers"]),
            ("Average order value", k["average_order_value"], d["average_order_value"]),
            ("Gross margin %", k["gross_margin_pct"], None),
            ("Return rate %", k["return_rate_pct"], None),
        ],
        columns=["Metric", "Value", "Change vs previous period %"],
    )
    rows = session.execute(
        select(
            S.order_id,
            S.order_date,
            S.customer_id,
            S.customer_name,
            S.product,
            S.category,
            S.region,
            S.channel,
            S.status,
            S.quantity,
            S.unit_price,
            S.discount,
            S.revenue,
            S.cost,
        )
        .where(*filters.conditions())
        .order_by(S.order_date, S.id)
        .limit(MAX_EXPORT_ROWS)
    ).all()
    columns = [
        "order_id", "order_date", "customer_id", "customer_name", "product", "category",
        "region", "channel", "status", "quantity", "unit_price", "discount", "revenue", "cost",
    ]  # fmt: skip
    raw = pd.DataFrame([tuple(r) for r in rows], columns=columns)
    for column in ("quantity", "unit_price", "discount", "revenue", "cost"):
        raw[column] = pd.to_numeric(raw[column])

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _sheet(writer, "Summary", summary)
        _sheet(writer, "Revenue over time", pd.DataFrame(overview["timeseries"]))
        for name, key in (
            ("By category", "by_category"),
            ("By region", "by_region"),
            ("By channel", "by_channel"),
            ("Top products", "top_products"),
            ("Top customers", "top_customers"),
        ):
            _sheet(writer, name, pd.DataFrame(overview[key], columns=["key", "revenue", "lines"]))
        if overview["anomalies"]:
            _sheet(writer, "Anomalies", pd.DataFrame(overview["anomalies"]))
        _sheet(writer, "Sales lines", raw)
    return buffer.getvalue()
